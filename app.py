import streamlit as st
import re
import datetime
import io
import base64
from weasyprint import HTML
from pptx import Presentation
import google.generativeai as genai
from PIL import Image

# ڕێکخستنی سەرەکی
st.set_page_config(page_title="داڕێژەری زیرەک", page_icon="✨", layout="wide")

if 'doc_content' not in st.session_state:
    st.session_state.doc_content = ""

st.markdown("""
<h1 style='text-align: center; color: #1e3a8a; font-weight: 900;'>داڕێژەری زیرەک</h1>
<p style='text-align: center; color: #64748b; font-size: 18px;'>پلاتفۆرمی بەرهەمهێنانی ناوەڕۆک (دەق و وێنە بۆ PDF و PowerPoint)</p>
<hr>
""", unsafe_allow_html=True)

tab_create, tab_image, tab_ai = st.tabs(["📝 داڕشتنی دەق (PDF / PPTX)", "🖼️ وێنە بۆ PDF", "🤖 مۆدی زیرەکی دەستکرد"])

with tab_create:
    col1, col2 = st.columns([2, 1])
    
    with col2:
        st.markdown("### 📌 زانیارییەکان")
        title = st.text_input("ناونیشانی بابەت:", "بابەتی نوێ")
        subtitle = st.text_input("ژێرنووس (پوختە):", "")
        
        st.markdown("---")
        export_format = st.radio("جۆری فایلەکە هەڵبژێرە:", ["دەمەوێت فایلی PDF دروست بکەم 📄", "دەمەوێت سلایدی پاوەرپۆینت دروست بکەم 📊"])
        
        audio_link = st.text_input("🔗 لینکی دەنگ/ڤیدیۆ (ئارەزوومەندانە):", placeholder="ئەگەر لینکت هەیە لێرە دایبنێ...")
        
        st.markdown("---")
        st.markdown("### 🎨 دیزاینەکان")
        active_theme = st.selectbox("قاڵبی پەڕەکە هەڵبژێرە:", [
            "دەفتەری تێبینی (هێڵکار)", 
            "پزیشکی و زانستی (شین)", 
            "کلاسیک و ئەدەبی (قاوەیی)",
            "تەکنەلۆژیا (تاریک و مۆدێرن)",
            "سروشتی و ژینگە (سەوز)",
            "شاهانە (ڕەش و زێڕین)",
            "ڕۆمانسی و شیعر (پەمەیی)",
            "فەرمی و ئیداری (ڕەساسی)",
            "تۆڕی زانستی (گرافیک)",
            "گفتوگۆ و زمان (مۆر)"
        ])
        
        cover_page = st.checkbox("پەڕەی بەرگ دروست بکە", value=True)
        show_flag = st.checkbox("ئاڵای کوردستان دابنێ", value=True)

    with col1:
        st.info("💡 **زانیاری:** دەتوانیت زمانی کوردی و ئینگلیزی تێکەڵ بکەیت، سیستەمەکە خۆی ئاڕاستەی وشەکان ڕاست دەکاتەوە.")
        raw_text = st.text_area("دەقەکەت لێرە دابنێ:", value=st.session_state.doc_content, height=400)
        
        if "PDF" in export_format:
            if st.button("📄 دروستکردنی فایلی PDF", use_container_width=True, type="primary"):
                if not raw_text.strip():
                    st.error("تکایە دەقێک بنووسە.")
                else:
                    with st.spinner("ئامادەکردنی PDF..."):
                        lines = [line.strip() for line in raw_text.split('\n') if line.strip()]
                        
                        # ١٠ دیزاینە جیاوازەکە
                        themes = {
                            "دەفتەری تێبینی (هێڵکار)": {"primary": "#334155", "bg": "repeating-linear-gradient(#f8fafc, #f8fafc 38px, #cbd5e1 38px, #cbd5e1 39px)", "card": "rgba(255,255,255,0.9)", "border": "#64748b", "accent": "#f1f5f9"},
                            "پزیشکی و زانستی (شین)": {"primary": "#0369a1", "bg": "#f0f9ff", "card": "#ffffff", "border": "#0ea5e9", "accent": "#bae6fd"},
                            "کلاسیک و ئەدەبی (قاوەیی)": {"primary": "#92400e", "bg": "#fefce8", "card": "#ffffff", "border": "#d97706", "accent": "#fde68a"},
                            "تەکنەلۆژیا (تاریک و مۆدێرن)": {"primary": "#38bdf8", "bg": "#0f172a", "card": "#1e293b", "border": "#0284c7", "accent": "#0f172a", "text": "#f8fafc"},
                            "سروشتی و ژینگە (سەوز)": {"primary": "#15803d", "bg": "#f0fdf4", "card": "#ffffff", "border": "#22c55e", "accent": "#bbf7d0"},
                            "شاهانە (ڕەش و زێڕین)": {"primary": "#fbbf24", "bg": "#171717", "card": "#262626", "border": "#d97706", "accent": "#171717", "text": "#fef3c7"},
                            "ڕۆمانسی و شیعر (پەمەیی)": {"primary": "#be185d", "bg": "#fdf2f8", "card": "#ffffff", "border": "#f43f5e", "accent": "#fbcfe8"},
                            "فەرمی و ئیداری (ڕەساسی)": {"primary": "#374151", "bg": "#f3f4f6", "card": "#ffffff", "border": "#9ca3af", "accent": "#e5e7eb"},
                            "تۆڕی زانستی (گرافیک)": {"primary": "#0f766e", "bg": "linear-gradient(#ccfbf1 1px, transparent 1px), linear-gradient(90deg, #ccfbf1 1px, transparent 1px)", "card": "rgba(255,255,255,0.95)", "border": "#14b8a6", "accent": "#ccfbf1"},
                            "گفتوگۆ و زمان (مۆر)": {"primary": "#6d28d9", "bg": "#f5f3ff", "card": "#ffffff", "border": "#8b5cf6", "accent": "#ddd6fe"}
                        }
                        
                        t_style = themes.get(active_theme)
                        text_color = t_style.get("text", "#0f172a")
                        
                        content_blocks = ""
                        for line in lines:
                            line = re.sub(r'\*\*(.*?)\*\*', f'<strong style="color: {t_style["primary"]};">\\1</strong>', line)
                            
                            # لێرە کێشەی دوو زمانی چارەسەر کراوە بە (dir="auto")
                            if line.startswith("#"):
                                content_blocks += f'<div dir="auto" style="text-align: start; background: {t_style["accent"]}; padding: 15px; border-radius: 8px; margin: 25px 0 15px 0; border-right: 6px solid {t_style["primary"]};"><h2 style="margin:0; color:{t_style["primary"]};">{line.replace("#", "").strip()}</h2></div>'
                            elif "|" in line:
                                cells = [c.strip() for c in line.split("|") if c.strip()]
                                row_html = "".join([f'<td style="border: 1px solid {t_style["border"]}; padding: 10px;">{c}</td>' for c in cells])
                                content_blocks += f'<table dir="auto" style="width:100%; border-collapse: collapse; margin: 15px 0; color: {text_color};"><tr>{row_html}</tr></table>'
                            else:
                                content_blocks += f'<div style="background:{t_style["card"]}; padding:15px; border-radius:8px; margin-bottom:10px; border-right:4px solid {t_style["border"]};"><p dir="auto" style="text-align: start; margin:0; color: {text_color};">{line}</p></div>'

                        audio_btn = f'<div style="text-align: center; margin: 30px 0;"><a href="{audio_link}" style="background-color: #e11d48; color: white; padding: 15px 30px; text-decoration: none; font-size: 16pt; border-radius: 10px; font-weight: bold; display: inline-block;">🔊 کلیک لێرە بکە بۆ کردنەوەی دەنگ / ڤیدیۆ</a></div>' if audio_link else ""

                        flag_html = ""
                        if show_flag:
                            flag_html = """<div style="position: absolute; top: -10px; left: 10px;"><svg viewBox="0 0 300 200" style="width: 50px; border-radius: 4px;"><rect width="300" height="66.6" fill="#ED2024"/><rect y="66.6" width="300" height="66.6" fill="#FFFFFF"/><rect y="133.2" width="300" height="66.8" fill="#278E43"/><circle cx="150" cy="100" r="22" fill="#FEBD11"/></svg></div>"""

                        cover_html = f'''
                        <div style="text-align: center; margin-top: 40%; page-break-after: always; position: relative;">
                            {flag_html}
                            <div style="display: inline-block; padding: 40px; border: 3px solid {t_style["border"]}; border-radius: 20px; background: {t_style["card"]};">
                                <h1 style="font-size: 250%; color: {t_style["primary"]}; border-bottom: 4px solid {t_style["accent"]}; padding-bottom: 20px;">{title}</h1>
                                <p style="font-size: 150%; color: {text_color};">{subtitle}</p>
                            </div>
                        </div>
                        ''' if cover_page else flag_html

                        bg_property = f'background: {t_style["bg"]};' if "gradient" in t_style["bg"] else f'background-color: {t_style["bg"]};'

                        final_html = f"""
                        <html>
                        <head><style>@page {{ size: A4; margin: 20mm; {bg_property} }} body {{ font-family: 'Amiri', Tahoma, sans-serif; font-size: 16pt; line-height: 1.8; }}</style></head>
                        <body>
                            {cover_html}
                            <h1 dir="auto" style="color:{t_style["primary"]}; text-align:center;">{title}</h1>
                            <h3 dir="auto" style="color:#64748b; text-align:center;">{subtitle}</h3>
                            {audio_btn}
                            {content_blocks}
                        </body></html>
                        """
                        
                        pdf_buf = io.BytesIO()
                        HTML(string=final_html).write_pdf(target=pdf_buf)
                        st.success("✅ فایلەکە بە سەرکەوتوویی ئامادە کرا.")
                        st.download_button("📥 داگرتنی فایلی PDF", data=pdf_buf.getvalue(), file_name=f"{title}.pdf", mime="application/pdf", use_container_width=True)

        elif "PowerPoint" in export_format:
            if st.button("📊 دروستکردنی سلایدی پاوەرپۆینت", use_container_width=True, type="primary"):
                if not raw_text.strip():
                    st.error("تکایە دەقێک بنووسە.")
                else:
                    with st.spinner("ئامادەکردنی پاوەرپۆینت..."):
                        prs = Presentation()
                        
                        # پەڕەی سەرەتا
                        title_slide_layout = prs.slide_layouts[0]
                        slide = prs.slides.add_slide(title_slide_layout)
                        slide.shapes.title.text = title
                        slide.placeholders[1].text = subtitle
                        if audio_link:
                            slide.placeholders[1].text += f"\n\nلینکی ڤیدیۆ/دەنگ: {audio_link}"
                        
                        lines = [line.strip() for line in raw_text.split('\n') if line.strip()]
                        current_slide = None
                        tf = None
                        
                        for line in lines:
                            clean_line = line.replace("**", "").replace("==", "")
                            if clean_line.startswith("#"):
                                slide_layout = prs.slide_layouts[1]
                                current_slide = prs.slides.add_slide(slide_layout)
                                current_slide.shapes.title.text = clean_line.replace("#", "").strip()
                                tf = current_slide.placeholders[1].text_frame
                            elif current_slide and tf:
                                p = tf.add_paragraph()
                                p.text = clean_line
                            else:
                                slide_layout = prs.slide_layouts[1]
                                current_slide = prs.slides.add_slide(slide_layout)
                                current_slide.shapes.title.text = "زانیاری گشتی"
                                tf = current_slide.placeholders[1].text_frame
                                p = tf.add_paragraph()
                                p.text = clean_line
                                
                        ppt_buf = io.BytesIO()
                        prs.save(ppt_buf)
                        st.success("✅ پاوەرپۆینتەکە ئامادەیە.")
                        st.download_button("📥 داگرتنی پاوەرپۆینت", data=ppt_buf.getvalue(), file_name=f"{title}.pptx", mime="application/vnd.openxmlformats-officedocument.presentationml.presentation", use_container_width=True)

with tab_image:
    st.markdown("### 🖼️ گۆڕینی وێنە بۆ PDF")
    st.info("ئەگەر تێبینی یان وێنەی کتێبت هەیە، لێرە ئەپڵۆدی بکە بۆ ئەوەی ڕاستەوخۆ بیکاتە فایلی PDF.")
    uploaded_images = st.file_uploader("وێنەکانت هەڵبژێرە:", type=['png', 'jpg', 'jpeg'], accept_multiple_files=True)
    
    if uploaded_images:
        if st.button("📄 دروستکردنی PDF لە وێنەکان", type="primary"):
            with st.spinner("خەریکی بەستنەوەی وێنەکانە..."):
                img_html_content = ""
                for img_file in uploaded_images:
                    encoded = base64.b64encode(img_file.read()).decode()
                    mime_type = img_file.type
                    img_html_content += f'<div style="text-align:center; page-break-after:always;"><img src="data:{mime_type};base64,{encoded}" style="max-width:100%; max-height:100vh; object-fit:contain;"></div>'
                
                final_img_pdf = f"<html><head><style>@page {{ size: A4; margin: 0; }} body {{ margin: 0; padding: 0; }}</style></head><body>{img_html_content}</body></html>"
                pdf_buf = io.BytesIO()
                HTML(string=final_img_pdf).write_pdf(target=pdf_buf)
                st.success("✅ وێنەکان کران بە فایلی PDF.")
                st.download_button("📥 داگرتنی PDF", data=pdf_buf.getvalue(), file_name="Images_to_Document.pdf", mime="application/pdf")

with tab_ai:
    st.markdown("### 🤖 دروستکردنی ناوەڕۆک بە زیرەکی دەستکرد")
    user_api_key = st.text_input("🔑 کلیلی تایبەتت (Gemini API Key):", type="password")
    
    ai_input_text = st.text_area("✍️ چی بنووسم بۆت؟", placeholder="نموونە: بابەتێک لەسەر چارەسەری نەخۆشی شەکرە...")
    
    if st.button("✨ داواکردن لە AI", type="primary"):
        if not user_api_key or not ai_input_text:
            st.error("کلیل و داواکارییەکە پڕ بکەرەوە.")
        else:
            with st.spinner("🤖 خەریکی نووسینە..."):
                try:
                    genai.configure(api_key=user_api_key)
                    model = genai.GenerativeModel('gemini-1.5-flash')
                    response = model.generate_content(f"وەڵامەکەت ڕێکبخە. بۆ هەر بەشێک سەردێڕی # بەکاربهێنە. \n\n {ai_input_text}")
                    st.session_state.doc_content = response.text
                    st.success("✅ ئامادەیە! بڕۆ تابـی یەکەم (داڕشتنی دەق) بۆ بینین و گۆڕینی بۆ PDF یان پاوەرپۆینت.")
                except Exception as e:
                    st.error(f"❌ هەڵەیەک ڕوویدا: {e}")

st.markdown("<br><hr><p style='text-align: center; color: #94a3b8;'>داڕێژەری زیرەک © 2026</p>", unsafe_allow_html=True)
